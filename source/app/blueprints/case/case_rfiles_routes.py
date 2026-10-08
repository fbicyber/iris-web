#  IRIS Source Code
#  Copyright (C) 2021 - Airbus CyberSecurity (SAS) - DFIR-IRIS Team
#  ir@cyberactionlab.net - contact@dfir-iris.org
#
#  This program is free software; you can redistribute it and/or
#  modify it under the terms of the GNU Lesser General Public
#  License as published by the Free Software Foundation; either
#  version 3 of the License, or (at your option) any later version.
#
#  This program is distributed in the hope that it will be useful,
#  but WITHOUT ANY WARRANTY; without even the implied warranty of
#  MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the GNU
#  Lesser General Public License for more details.
#
#  You should have received a copy of the GNU Lesser General Public License
#  along with this program; if not, write to the Free Software Foundation,
#  Inc., 51 Franklin Street, Fifth Floor, Boston, MA  02110-1301, USA.

# IMPORTS ------------------------------------------------
from datetime import datetime

import marshmallow
from flask import Blueprint
from flask import redirect
from flask import render_template
from flask import request
from flask import url_for
from flask_login import current_user
from flask_wtf import FlaskForm
from openpyxl import Workbook, load_workbook
from io import BytesIO

from app.extensions import db
from flask import current_app as app
from app.blueprints.case.case_comments import case_comment_update
from app.datamgmt.case.case_db import get_case
from app.datamgmt.case.case_rfiles_db import add_comment_to_evidence
from app.datamgmt.case.case_rfiles_db import add_rfile
from app.datamgmt.case.case_rfiles_db import delete_evidence_comment
from app.datamgmt.case.case_rfiles_db import delete_rfile
from app.datamgmt.case.case_rfiles_db import get_case_evidence_comment
from app.datamgmt.case.case_rfiles_db import get_case_evidence_comments
from app.datamgmt.case.case_rfiles_db import get_case_evidence_comments_count
from app.datamgmt.case.case_rfiles_db import get_rfile, get_rfile_from_ext_id
from app.datamgmt.case.case_rfiles_db import get_rfiles
from app.datamgmt.case.case_rfiles_db import update_rfile
from app.datamgmt.case.case_rfiles_db import get_evidence_type_by_name, get_default_evidence_type
from app.datamgmt.manage.manage_attribute_db import get_default_custom_attributes
from app.datamgmt.states import get_evidences_state
from app.iris_engine.module_handler.module_handler import call_modules_hook
from app.iris_engine.utils.tracker import track_activity
from app.models.authorization import CaseAccessLevel
from app.schema.marshables import CaseEvidenceSchema
from app.schema.marshables import CommentSchema
from app.util import ac_api_case_requires
from app.util import ac_case_requires
from app.util import response_error
from app.util import response_success

case_rfiles_blueprint = Blueprint(
    'case_rfiles',
    __name__,
    template_folder='templates'
)


# CONTENT ------------------------------------------------
@case_rfiles_blueprint.route('/case/evidences', methods=['GET'])
@ac_case_requires(CaseAccessLevel.read_only, CaseAccessLevel.full_access)
def case_rfile(caseid, url_redir):
    if url_redir:
        return redirect(url_for('case_rfiles.case_rfile', cid=caseid, redirect=True))

    form = FlaskForm()
    case = get_case(caseid)

    return render_template("case_rfile.html", case=case, form=form)


@case_rfiles_blueprint.route('/case/evidences/list', methods=['GET'])
@ac_api_case_requires(CaseAccessLevel.read_only, CaseAccessLevel.full_access)
def case_list_rfiles(caseid):
    crf = get_rfiles(caseid)

    ret = {
        "evidences": CaseEvidenceSchema().dump(crf, many=True),
        "state": get_evidences_state(caseid=caseid)
    }

    return response_success("", data=ret)


@case_rfiles_blueprint.route('/case/evidences/state', methods=['GET'])
@ac_api_case_requires(CaseAccessLevel.read_only, CaseAccessLevel.full_access)
def case_rfiles_state(caseid):
    os = get_evidences_state(caseid=caseid)
    if os:
        return response_success(data=os)
    else:
        return response_error('No evidences state for this case.')


@case_rfiles_blueprint.route('/case/evidences/add', methods=['POST'])
@ac_api_case_requires(CaseAccessLevel.full_access)
def case_add_rfile(caseid):

    try:
        # validate before saving
        evidence_schema = CaseEvidenceSchema()

        request_data = call_modules_hook('on_preload_evidence_create', data=request.get_json(), caseid=caseid)
        # Never load a client-supplied primary key on create, or marshmallow-sqlalchemy
        # would fetch and overwrite an existing evidence record from another case.
        request_data.pop('id', None)

        evidence = evidence_schema.load(request_data)

        crf = add_rfile(evidence=evidence,
                        user_id=current_user.id,
                        caseid=caseid
                         )

        crf = call_modules_hook('on_postload_evidence_create', data=crf, caseid=caseid)

        if crf:
            track_activity(f"added evidence \"{crf.filename}\"", caseid=caseid)
            return response_success("Evidence added", data=evidence_schema.dump(crf))

        return response_error("Unable to create task for internal reasons")

    except marshmallow.exceptions.ValidationError as e:
        return response_error(msg="Data error", data=e.messages)


@case_rfiles_blueprint.route('/case/evidences/<int:cur_id>', methods=['GET'])
@ac_api_case_requires(CaseAccessLevel.read_only, CaseAccessLevel.full_access)
def case_get_evidence(cur_id, caseid):
    crf = get_rfile(cur_id, caseid)
    if not crf:
        return response_error("Invalid evidence ID for this case")

    evidence_schema = CaseEvidenceSchema()
    return response_success(data=evidence_schema.dump(crf))


@case_rfiles_blueprint.route('/case/evidences/<int:cur_id>/modal', methods=['GET'])
@ac_case_requires(CaseAccessLevel.read_only, CaseAccessLevel.full_access)
def case_edit_rfile_modal(cur_id, caseid, url_redir):
    if url_redir:
        return redirect(url_for('case_rfiles.case_rfile', cid=caseid, redirect=True))

    crf = get_rfile(cur_id, caseid)
    if not crf:
        return response_error("Invalid evidence ID for this case")

    comments_map = get_case_evidence_comments_count([cur_id])

    return render_template("modal_add_case_rfile.html", rfile=crf, attributes=crf.custom_attributes,
                           comments_map=comments_map)


@case_rfiles_blueprint.route('/case/evidences/add/modal', methods=['GET'])
@ac_api_case_requires(CaseAccessLevel.full_access)
def case_add_rfile_modal(caseid):

    return render_template("modal_add_case_rfile.html", rfile=None, attributes=get_default_custom_attributes('evidence'))


@case_rfiles_blueprint.route('/case/evidences/update/<int:cur_id>', methods=['POST'])
@ac_api_case_requires(CaseAccessLevel.full_access)
def case_edit_rfile(cur_id, caseid):

    try:
        # validate before saving
        evidence_schema = CaseEvidenceSchema()

        request_data = call_modules_hook('on_preload_evidence_update', data=request.get_json(), caseid=caseid)

        crf = get_rfile(cur_id, caseid)
        if not crf:
            return response_error("Invalid evidence ID for this case")

        request_data['id'] = cur_id
        evidence = evidence_schema.load(request_data, instance=crf)

        evd = update_rfile(evidence=evidence,
                           user_id=current_user.id,
                           caseid=caseid
                           )

        evd = call_modules_hook('on_postload_evidence_update', data=evd, caseid=caseid)

        if evd:
            track_activity(f"updated evidence \"{evd.filename}\"", caseid=caseid)
            return response_success("Evidence {} updated".format(evd.filename), data=evidence_schema.dump(evd))

        return response_error("Unable to update task for internal reasons")

    except marshmallow.exceptions.ValidationError as e:
        return response_error(msg="Data error", data=e.messages)


@case_rfiles_blueprint.route('/case/evidences/delete/<int:cur_id>', methods=['POST'])
@ac_api_case_requires(CaseAccessLevel.full_access)
def case_delete_rfile(cur_id, caseid):

    call_modules_hook('on_preload_evidence_delete', data=cur_id, caseid=caseid)
    crf = get_rfile(cur_id, caseid)
    if not crf:
        return response_error("Invalid evidence ID for this case")

    delete_rfile(cur_id, caseid=caseid)

    call_modules_hook('on_postload_evidence_delete', data=cur_id, caseid=caseid)

    track_activity(f"deleted evidence \"{crf.filename}\" from registry", caseid)

    return response_success("Evidence deleted")


@case_rfiles_blueprint.route('/case/evidences/<int:cur_id>/comments/modal', methods=['GET'])
@ac_case_requires(CaseAccessLevel.read_only, CaseAccessLevel.full_access)
def case_comment_evidence_modal(cur_id, caseid, url_redir):
    if url_redir:
        return redirect(url_for('case_task.case_task', cid=caseid, redirect=True))

    evidence = get_rfile(cur_id, caseid=caseid)
    if not evidence:
        return response_error('Invalid evidence ID')

    return render_template("modal_conversation.html", element_id=cur_id, element_type='evidences',
                           title=evidence.filename)


@case_rfiles_blueprint.route('/case/evidences/<int:cur_id>/comments/list', methods=['GET'])
@ac_api_case_requires(CaseAccessLevel.read_only, CaseAccessLevel.full_access)
def case_comment_evidence_list(cur_id, caseid):

    evidence_comments = get_case_evidence_comments(cur_id, caseid)
    if evidence_comments is None:
        return response_error('Invalid evidence ID')

    return response_success(data=CommentSchema(many=True).dump(evidence_comments))


@case_rfiles_blueprint.route('/case/evidences/<int:cur_id>/comments/add', methods=['POST'])
@ac_api_case_requires(CaseAccessLevel.full_access)
def case_comment_evidence_add(cur_id, caseid):

    try:
        evidence = get_rfile(cur_id, caseid=caseid)
        if not evidence:
            return response_error('Invalid evidence ID')

        comment_schema = CommentSchema()

        comment = comment_schema.load(request.get_json())
        comment.comment_case_id = caseid
        comment.comment_user_id = current_user.id
        comment.comment_date = datetime.now()
        comment.comment_update_date = datetime.now()
        db.session.add(comment)
        db.session.commit()

        add_comment_to_evidence(evidence.id, comment.comment_id)

        db.session.commit()

        hook_data = {
            "comment": comment_schema.dump(comment),
            "evidence": CaseEvidenceSchema().dump(evidence)
        }
        call_modules_hook('on_postload_evidence_commented', data=hook_data, caseid=caseid)

        track_activity(f"evidence \"{evidence.filename}\" commented", caseid=caseid)
        return response_success("Event commented", data=comment_schema.dump(comment))

    except marshmallow.exceptions.ValidationError as e:
        return response_error(msg="Data error", data=e.normalized_messages())


@case_rfiles_blueprint.route('/case/evidences/<int:cur_id>/comments/<int:com_id>', methods=['GET'])
@ac_api_case_requires(CaseAccessLevel.read_only, CaseAccessLevel.full_access)
def case_comment_evidence_get(cur_id, com_id, caseid):

    comment = get_case_evidence_comment(cur_id, com_id, caseid)
    if not comment:
        return response_error("Invalid comment ID")

    return response_success(data=comment._asdict())


@case_rfiles_blueprint.route('/case/evidences/<int:cur_id>/comments/<int:com_id>/edit', methods=['POST'])
@ac_api_case_requires(CaseAccessLevel.full_access)
def case_comment_evidence_edit(cur_id, com_id, caseid):

    return case_comment_update(com_id, 'tasks', caseid)


@case_rfiles_blueprint.route('/case/evidences/<int:cur_id>/comments/<int:com_id>/delete', methods=['POST'])
@ac_api_case_requires(CaseAccessLevel.full_access)
def case_comment_evidence_delete(cur_id, com_id, caseid):

    success, msg = delete_evidence_comment(cur_id, com_id)
    if not success:
        return response_error(msg)

    call_modules_hook('on_postload_evidence_comment_delete', data=com_id, caseid=caseid)

    track_activity(f"comment {com_id} on evidence {cur_id} deleted", caseid=caseid)
    return response_success(msg)


@case_rfiles_blueprint.route('/case/evidences/excel_upload', methods=['POST'])
@ac_api_case_requires(CaseAccessLevel.full_access)
def case_evidences_upload_excel(caseid):
    # to validate uploaded evidence attributes
    evidence_schema = CaseEvidenceSchema()

    jsdata = request.get_json()
    if not jsdata or "excel_data" not in jsdata:
        return response_error(msg="Unable to get data imported from Excel", data={"Exception": f"Unable to get data imported from Excel"})

    app.logger.info("Starting Excel import")
    evidence_fields = [
        "id",
        "filename",
        "type",
        "file_hash",
        "file_size",
        "file_description",
        "host",
        "external_id"
    ]

    list_of_errors = []

    # excel data is received as an array of numbers, actually uint8 converted by js
    excel_data_bytes = jsdata["excel_data"]

    # convert the array of numbers to a byte array, then a bytestring, then a mock(?) file object for load_workbook to read
    workbook = load_workbook(BytesIO(bytes(bytearray(excel_data_bytes))))
    worksheet = workbook.worksheets[0]
    excel_lines = []
    for i, row in enumerate(worksheet):
        if i == 0:
            headers = [cell.value for cell in row]
            missing_fields = [fld for fld in evidence_fields if fld not in headers]
            if len(missing_fields) > 0:
                msg = f"Bad XLSX Fields Mapping. Fields missing: [{','.join(missing_fields)}]"
                data = {"error_code": "BAD_FIELDS_MAPPING", "expected": ','.join(evidence_fields), "found": ','.join(headers),
                        "missing": ','.join(missing_fields)}
                app.logger.warning(data)

                return response_error(msg=msg, data=data)
        else:
            line = []
            for i, cell in enumerate(row):
                line.append(cell.value)
            excel_lines.append(line)

    DEFAULT_EVIDENCE_TYPE_ID = get_default_evidence_type().id
    import_error = False
    # ==========================  checking data validity  ==========================
    row_index = 1
    excel_lines_to_save = []
    for row in excel_lines:
        try:
            row_index += 1
            if not any(row):
                continue
            row_to_save = {}

            # get attributes from each row
            filename = str(row[headers.index('filename')])
            evidence_type = row[headers.index('type')]            
            file_hash = str(row[headers.index('file_hash')])
            file_size = str(row[headers.index('file_size')])
            file_description = str(row[headers.index('file_description')])
            host = str(row[headers.index('host')])
            external_id = str(row[headers.index('external_id')])

            # check for evidence id
            if row[headers.index('id')]:
                row_to_save['id'] = row[headers.index('id')]
            else:
                # check for any external id
                if row[headers.index('external_id')]:
                    crf = get_rfile_from_ext_id(external_id, caseid)
                    if crf:
                        row_to_save['id'] = crf.id


            # error checking filename
            if filename is None or len(filename) == 0:
                error_msg = f"Recoverable in row {row_index}, Evidence Filename cannot be empty."
                app.logger.error(error_msg)
                list_of_errors.append(error_msg)
                import_error = True
            else:
                row_to_save['filename'] = filename
            
            # error checking evidence type
            if (evidence_type != None) and (evidence_type != ""):
                try:
                    valid_evidence_type = get_evidence_type_by_name(evidence_type)
                    row_to_save['type_id'] = valid_evidence_type.id

                except Exception as e:
                    error_msg = f"Recoverable in row {row_index}, evidence type not recognized: {evidence_type}."
                    app.logger.error(error_msg)
                    list_of_errors.append(error_msg)
                    import_error = True
                    row_to_save['type_id'] = DEFAULT_EVIDENCE_TYPE_ID
            
            # adding file size, hash, description, additional metadata
            row_to_save['file_size'] = file_size
            row_to_save['file_hash'] = file_hash
            row_to_save['file_description'] = file_description
            row_to_save['host'] = host
            row_to_save['external_id'] = external_id
            
            # appending row
            app.logger.info(f"Appending row {row_index}")
            excel_lines_to_save.append(row_to_save)

        except Exception as e:
            return response_error(msg=f"Data error", data={"Exception": f"Unhandled error {e}.\nrow number: {row_index}"})
    
    # ========================== begin saving data ============================
    session = db.session.begin_nested()
    row_index = 1
    for row in excel_lines_to_save:
        if row is None:
            continue
        row_index += 1
        app.logger.info(f"Saving ROW {row_index}")

        try:
            if "id" in row:
                request_data = call_modules_hook('on_preload_evidence_update', data=row, caseid=caseid)
            else:
                request_data = call_modules_hook('on_preload_evidence_create', data=row, caseid=caseid)
            evidence = evidence_schema.load(request_data)

            # add evidence if no id exists
            if "id" not in row:
                crf = add_rfile(evidence=evidence,
                    user_id=current_user.id,
                    caseid=caseid
                )
                crf = call_modules_hook('on_postload_evidence_create', data=crf, caseid=caseid)
            
                if crf:
                    track_activity(f"Added evidence \"{crf.filename}\"", caseid=caseid)
                    app.logger.info(f"Evidence added: {crf.filename}")

            else:
                # update evidence if id exists
                crf = get_rfile(row["id"], caseid)
                evd = update_rfile(evidence=evidence,
                    user_id=current_user.id,
                    caseid=caseid
                )

                evd = call_modules_hook('on_postload_evidence_update', data=evd, caseid=caseid)
                if evd:
                    track_activity(f"updated evidence \"{evd.filename}\"", caseid=caseid)
                    app.logger.info(f"Evidence updated: {crf.filename}")

        except marshmallow.exceptions.ValidationError as e:
            error_msg = f"Unrecoverable error in row {row_index} while validating, Exception : {e}"
            app.logger.error(error_msg)
            list_of_errors.append(error_msg)
            import_error = True

        except Exception as e:
            error_msg = f"Unrecoverable error in row {row_index} at unknown point, Exception : {e}"
            app.logger.error(error_msg)
            list_of_errors.append(error_msg)
            import_error = True    
    
    try:
        session.commit()
    except:
        pass
        
    app.logger.info("======================== END_EXCEL_IMPORT ==========================================")
    if not import_error:
        return response_success(msg="Evidence added with no errors (Excel File)")
    else:
        return response_success(msg=f"Events added with errors: {list_of_errors}", data=list_of_errors)
